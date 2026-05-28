import openpyxl
from pathlib import Path
from langchain_core.documents import Document

def load_excel():
    excel_path = Path(__file__).parent / "data" / "test_cases.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    docs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        record = dict(zip(headers, row))
        content = (
            f"Module: {record.get('Module Name', '')}\n"
            f"Feature: {record.get('Feature', '')}\n"
            f"Test Scenario: {record.get('Test Scenario', '')}\n"
            f"Test Steps: {record.get('Test Steps', '')}\n"
            f"Expected Result: {record.get('Expected Result', '')}\n"
            f"Priority: {record.get('Priority', '')}\n"
            f"Regression Category: {record.get('Regression Category', '')}\n"
            f"Test Case ID: {record.get('Test Case ID', '')}"
        )
        docs.append(Document(page_content=content, metadata={"Test Case ID": record.get("Test Case ID", ""), "Module": record.get("Module Name", ""), "Feature": record.get("Feature", "")}))
    return docs

def chunk_documents(docs, chunk_size=500, overlap=50):
    chunks = []
    for doc in docs:
        content = doc.page_content
        start = 0
        while start < len(content):
            end = start + chunk_size
            chunk_text = content[start:end]
            chunk_meta = doc.metadata.copy()
            chunk_meta["chunk_index"] = len(chunks)
            chunks.append(Document(page_content=chunk_text, metadata=chunk_meta))
            start += chunk_size - overlap
    return chunks